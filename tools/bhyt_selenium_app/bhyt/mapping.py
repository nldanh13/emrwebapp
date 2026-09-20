from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DOC_BHXH = "BHXH07"
DOC_GRV = "GRV03"


DEFAULT_DOCTORS = [
    "Nguyễn Lê Hoan",
    "Nguyễn Chí Nguyện",
    "Phạm Việt Tân",
    "Phan Văn Tuấn",
    "Trần Quốc Toản",
    "Phạm Việt Triều",
    "Hoàng Minh Tú",
    "Trần Quang Sơn",
    "Bùi Công Minh",
    "Hồ Điền",
    "Vũ Tấn Thọ",
    "Đặng Phước Giàu",
    "Nguyễn Lâm Minh Tân",
    "Trần Nguyễn Anh Duy",
    "Nguyễn Tư Thái Bảo",
    "Nguyễn Giang Tử",
]


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("Đ", "D").replace("đ", "d").lower()
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    prefixes = {"ths", "th", "ts", "bs", "cki", "ckii", "bscki", "bsckii"}
    parts = text.split()
    while parts and parts[0] in prefixes:
        parts.pop(0)
    return " ".join(parts)


def header_key(value: Any) -> str:
    return normalize_text(value)


def cell_text(cell: Any) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%d/%m/%Y %H:%M")
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, int):
        fmt = str(cell.number_format or "")
        if re.fullmatch(r"0+", fmt):
            return f"{value:0{len(fmt)}d}"
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


@dataclass
class Record:
    doc_type: str
    source_file: str
    source_sheet: str
    source_row: int
    fields: dict[str, str]
    raw: dict[str, str]
    issues: list[str] = field(default_factory=list)

    @property
    def patient_name(self) -> str:
        return self.fields.get("ho_ten", "")

    @property
    def doctor_name(self) -> str:
        return self.fields.get("doctor_text", "")

    @property
    def record_key(self) -> str:
        if self.doc_type == DOC_BHXH:
            parts = [
                self.fields.get("so_seri", ""),
                self.fields.get("ma_bhxh", ""),
                self.fields.get("ho_ten", ""),
                self.fields.get("tu_ngay", ""),
                self.fields.get("den_ngay", ""),
            ]
        else:
            parts = [
                self.fields.get("so_seri", ""),
                self.fields.get("ma_ct", ""),
                self.fields.get("ma_bhxh", ""),
                self.fields.get("ho_ten", ""),
                self.fields.get("tu_ngay", ""),
                self.fields.get("den_ngay", ""),
            ]
        payload = self.doc_type + "|" + "|".join(normalize_text(x) for x in parts)
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()

    def to_store(self) -> dict[str, Any]:
        return {
            "record_key": self.record_key,
            "doc_type": self.doc_type,
            "source_file": self.source_file,
            "source_sheet": self.source_sheet,
            "source_row": self.source_row,
            "patient_name": self.patient_name,
            "doctor_name": self.doctor_name,
            "fields_json": json.dumps(self.fields, ensure_ascii=False),
            "raw_json": json.dumps(self.raw, ensure_ascii=False),
            "issues_json": json.dumps(self.issues, ensure_ascii=False),
        }


def _value(row: dict[str, str], *aliases: str) -> str:
    normalized = {header_key(key): value for key, value in row.items()}
    for alias in aliases:
        value = normalized.get(header_key(alias), "")
        if value != "":
            return value
    return ""


def _find_header_row(ws: Any) -> tuple[int, list[str]] | None:
    for row_index in range(1, min(ws.max_row, 20) + 1):
        headers = [cell_text(cell) for cell in ws[row_index]]
        keys = {header_key(item) for item in headers if item}
        if "ho ten" in keys and (
            "nguoi hanh nghe" in keys or "truong khoa" in keys
        ):
            return row_index, headers
    return None


def _detect_type(headers: Iterable[str]) -> str | None:
    keys = {header_key(item) for item in headers if item}
    if "nguoi hanh nghe" in keys and (
        "dieu tri tu ngay" in keys or "so seri" in keys
    ):
        return DOC_BHXH
    if "truong khoa" in keys and (
        "ngay vao vien" in keys or "ma y te" in keys
    ):
        return DOC_GRV
    return None


def _map_bhxh(row: dict[str, str]) -> dict[str, str]:
    return {
        "ma_ct": _value(row, "Mã chứng từ"),
        "so_seri": _value(row, "Số seri"),
        "so_kcb": _value(row, "Số KCB"),
        "mau_so": _value(row, "Mẫu số") or "07",
        "ma_bhxh": _value(row, "Mã số BH", "Mã số BHXH"),
        "ma_the": _value(row, "Mã thẻ"),
        "ho_ten": _value(row, "Họ tên"),
        "ngay_sinh": _value(row, "Ngày sinh"),
        "gioi_tinh": _value(row, "Giới tính"),
        "ten_dv": _value(row, "Đơn vị"),
        "ngay_kcb": _value(row, "Ngày khám bệnh", "Ngày KCB"),
        "chan_doan": _value(row, "Chẩn đoán"),
        "tu_ngay": _value(row, "Điều trị từ ngày", "Từ ngày"),
        "den_ngay": _value(row, "Điều trị đến ngày", "Đến ngày"),
        "ho_ten_cha": _value(row, "Họ tên cha"),
        "ho_ten_me": _value(row, "Họ tên mẹ", "Họ tên mẹ / Người nuôi dưỡng"),
        "nguoi_dai_dien": _value(row, "Thủ trưởng", "Thủ trưởng đơn vị"),
        "doctor_text": _value(row, "Bác sĩ trong danh sách", "Người hành nghề"),
        "ngay_ct": _value(row, "Ngày chứng từ"),
        "loai_giay_to": _value(row, "Loại giấy tờ"),
        "so_cccd": _value(row, "Số CCCD", "CCCD"),
        "ngaycap_cccd": _value(row, "Ngày cấp CCCD", "Ngày cấp"),
        "noicap_cccd": _value(row, "Nơi cấp CCCD", "Nơi cấp"),
    }


def _map_grv(row: dict[str, str]) -> dict[str, str]:
    return {
        "ma_ct": _value(row, "Số lưu trữ"),
        "so_seri": _value(row, "Mã y tế"),
        "ma_bhxh": _value(row, "Mã số BH", "Mã số BHXH"),
        "ma_the": _value(row, "Mã thẻ"),
        "ho_ten": _value(row, "Họ tên"),
        "ngay_sinh": _value(row, "Ngày sinh"),
        "ho_ten_me": _value(row, "Họ tên mẹ / Người nuôi dưỡng", "Họ tên mẹ"),
        "ho_ten_cha": _value(row, "Họ tên cha"),
        "gioi_tinh": _value(row, "Giới tính"),
        "nghe_nghiep": _value(row, "Nghề nghiệp"),
        "ma_khoa": _value(row, "Khoa"),
        "dan_toc": _value(row, "Dân tộc"),
        "dia_chi": _value(row, "Địa chỉ"),
        "tu_ngay": _value(row, "Ngày vào viện"),
        "den_ngay": _value(row, "Ngày ra viện"),
        "dc_thainghen": _value(row, "Đình chỉ thai nghén"),
        "tuoi_thai": _value(row, "Tuổi thai"),
        "chan_doan": _value(row, "Chẩn đoán"),
        "pp_dieutri": _value(row, "Phương pháp điều trị"),
        "ghi_chu": _value(row, "Ghi chú"),
        "nguoi_dai_dien": _value(row, "Thủ trưởng đơn vị"),
        "doctor_text": _value(row, "Bác sĩ trong danh sách", "Trưởng khoa"),
        "ngay_ct": _value(row, "Ngày chứng từ"),
        "ngoaitru_tungay": _value(row, "Điều trị ngoại trú từ ngày"),
        "ngoaitru_denngay": _value(row, "Điều trị ngoại trú đến ngày"),
        "loai_giay_to": _value(row, "Loại giấy tờ") or "Không có giấy tờ",
        "so_cccd": _value(row, "Số CCCD", "CCCD"),
        "ngaycap_cccd": _value(row, "Ngày cấp CCCD", "Ngày cấp"),
        "noicap_cccd": _value(row, "Nơi cấp CCCD", "Nơi cấp"),
    }


def validate_fields(doc_type: str, fields: dict[str, str]) -> list[str]:
    if doc_type == DOC_BHXH:
        required = {
            "so_kcb": "Thiếu Số KCB (file xuất không có cột này)",
            "ma_bhxh": "Thiếu mã số BHXH",
            "ho_ten": "Thiếu họ tên",
            "ngay_sinh": "Thiếu ngày sinh",
            "chan_doan": "Thiếu chẩn đoán và điều trị",
            "tu_ngay": "Thiếu ngày bắt đầu nghỉ",
            "den_ngay": "Thiếu ngày kết thúc nghỉ",
            "doctor_text": "Thiếu người hành nghề",
            "ngay_ct": "Thiếu ngày chứng từ",
        }
    else:
        required = {
            "ma_bhxh": "Thiếu mã số BHXH",
            "ho_ten": "Thiếu họ tên",
            "ngay_sinh": "Thiếu ngày sinh",
            "gioi_tinh": "Thiếu giới tính",
            "ma_khoa": "Thiếu khoa",
            "dan_toc": "Thiếu dân tộc",
            "dia_chi": "Thiếu địa chỉ",
            "tu_ngay": "Thiếu ngày vào viện",
            "den_ngay": "Thiếu ngày ra viện",
            "chan_doan": "Thiếu chẩn đoán",
            "pp_dieutri": "Thiếu phương pháp điều trị",
            "doctor_text": "Thiếu người hành nghề/trưởng khoa",
            "ngay_ct": "Thiếu ngày chứng từ",
        }
    return [message for key, message in required.items() if not fields.get(key, "").strip()]


def load_records(paths: Iterable[str | Path], allowed_doctors: Iterable[str] | None = None) -> list[Record]:
    allowed = {normalize_text(name) for name in (allowed_doctors or DEFAULT_DOCTORS)}
    records: list[Record] = []
    seen_keys: set[str] = set()

    for path_like in paths:
        path = Path(path_like)
        workbook = load_workbook(path, data_only=True, read_only=False)
        for ws in workbook.worksheets:
            header_info = _find_header_row(ws)
            if not header_info:
                continue
            header_row, headers = header_info
            doc_type = _detect_type(headers)
            if not doc_type:
                continue
            width = len(headers)
            for row_index in range(header_row + 1, ws.max_row + 1):
                cells = list(ws[row_index])[:width]
                values = [cell_text(cell) for cell in cells]
                if not any(values):
                    continue
                raw = {headers[i] or f"Cột {i + 1}": values[i] for i in range(width)}
                fields = _map_bhxh(raw) if doc_type == DOC_BHXH else _map_grv(raw)
                doctor = normalize_text(fields.get("doctor_text"))
                if allowed and doctor not in allowed:
                    continue
                source_row = row_index
                try:
                    source_row = int(_value(raw, "Dòng nguồn") or row_index)
                except ValueError:
                    pass
                record = Record(
                    doc_type=doc_type,
                    source_file=path.name,
                    source_sheet=ws.title,
                    source_row=source_row,
                    fields=fields,
                    raw=raw,
                )
                record.issues = validate_fields(doc_type, fields)
                if record.record_key in seen_keys:
                    continue
                seen_keys.add(record.record_key)
                records.append(record)
    return records
